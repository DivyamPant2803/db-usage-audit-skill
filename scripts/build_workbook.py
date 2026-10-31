#!/usr/bin/env python3
"""Convert DB_USAGE_REPORT.md (or a docs/db-audit/ folder) into an Excel workbook.

Reads every markdown pipe table in the report, recognises each by its header
columns (see references/report-template.md), and writes one sheet per kind:

  Columns, Table_Indexes, Foreign_Keys, Index_Inventory, Workflows,
  Workflow_Access, Access_Flat, Table_to_Workflows, Column_to_Workflows,
  Unresolved

Access_Flat is the pivot-friendly sheet: one row per
(workflow, table, role, column) with its index-coverage label, built by
splitting the comma-separated column cells of Workflow_Access.

Usage: build_workbook.py <report.md | report-dir> [--out DB_USAGE_REPORT.xlsx]
Requires: openpyxl (pip install openpyxl). Python 3.8+.
"""
import argparse
import glob
import os
import re
import sys

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.exit("error: openpyxl is required. Run: pip install openpyxl")

CELL_LIMIT = 32000  # Excel hard limit is 32767 chars per cell
ROLE_COLUMNS = [
    ("Columns read", "READ"),
    ("Filter columns", "FILTER"),
    ("Join columns", "JOIN"),
    ("Sort/group", "SORT_GROUP"),
    ("Columns written", "WRITE"),
]
EMPTY_MARKERS = {"", "-", "—", "n/a", "none", "none found"}

H_TABLE = re.compile(r"^#{2,6}\s+`?(T-[^`\s(]+)`?")
H_WORKFLOW = re.compile(r"^#{2,6}\s+`?(W-\d+)`?\s*(.*)")


def split_row(line):
    """Split a markdown table row on unescaped pipes."""
    s = line.strip()
    if s.startswith("|"):
        s = s[1:]
    if s.endswith("|") and not s.endswith("\\|"):
        s = s[:-1]
    cells = re.split(r"(?<!\\)\|", s)
    return [clean(c) for c in cells]


def clean(text):
    t = text.strip().replace("\\|", "|").replace("`", "")
    t = re.sub(r"<br\s*/?>", "\n", t, flags=re.I)
    return t[:CELL_LIMIT]


def is_separator(line):
    return bool(re.match(r"^\s*\|?\s*:?-{2,}:?\s*(\|\s*:?-{2,}:?\s*)*\|?\s*$", line))


def parse_tables(text):
    """Yield (headers, rows, ctx) for each pipe table, tracking T-/W- headings."""
    lines = text.splitlines()
    ctx = {"table": "", "workflow": "", "workflow_title": ""}
    in_code = False
    i = 0
    while i < len(lines):
        line = lines[i]
        if line.strip().startswith("```"):
            in_code = not in_code
            i += 1
            continue
        if in_code:
            i += 1
            continue
        m = H_TABLE.match(line)
        if m:
            ctx["table"] = m.group(1)[2:]  # drop "T-"
        m = H_WORKFLOW.match(line)
        if m:
            ctx["workflow"] = m.group(1)
            ctx["workflow_title"] = clean(m.group(2))
        if line.lstrip().startswith("|") and i + 1 < len(lines) and is_separator(lines[i + 1]):
            headers = split_row(line)
            rows = []
            j = i + 2
            while j < len(lines) and lines[j].lstrip().startswith("|"):
                row = split_row(lines[j])
                row += [""] * (len(headers) - len(row))
                rows.append(row[: len(headers)])
                j += 1
            yield headers, rows, dict(ctx)
            i = j
            continue
        i += 1


def classify(headers):
    h = [x.lower() for x in headers]
    has = lambda name: any(x.startswith(name.lower()) for x in h)
    if has("table.column"):
        return "Column_to_Workflows"
    if has("column") and has("sql type"):
        return "Columns"
    if has("column") and has("references"):
        return "Foreign_Keys"
    if has("index id") and has("table"):
        return "Index_Inventory"
    if has("index id") and has("key columns"):
        return "Table_Indexes"
    if has("id") and has("trigger"):
        return "Workflows"
    if has("table") and has("op") and has("index coverage"):
        return "Workflow_Access"
    if has("table") and has("reads"):
        return "Table_to_Workflows"
    if has("item") and has("location"):
        return "Unresolved"
    return None


def split_columns(cell):
    """Split a cell like 'Id, Total (sum), Customer.Name' into column names."""
    if cell.strip().lower() in EMPTY_MARKERS:
        return []
    parts = re.split(r"[,;\n]", cell)
    out = []
    for p in parts:
        name = re.sub(r"\s*\(.*?\)\s*$", "", p.strip())
        if name and name.lower() not in EMPTY_MARKERS:
            out.append(name)
    return out


def parse_coverage(cell):
    """'CustomerId: COVERED-SEEK (IX-A); Status: NOT-COVERED' -> {col: (label, index)}"""
    result = {}
    for part in re.split(r"[;\n]", cell):
        m = re.match(r"\s*([^:]+):\s*([A-Z\-]+)\s*(?:\(([^)]*)\))?", part)
        if m:
            result[m.group(1).strip()] = (m.group(2), (m.group(3) or "").strip())
    return result


def build_flat(access_headers, access_rows):
    idx = {h.lower(): n for n, h in enumerate(access_headers)}
    get = lambda row, name: row[idx[name.lower()]] if name.lower() in idx else ""
    flat = []
    for row in access_rows:
        cov = parse_coverage(get(row, "Index coverage"))
        for col_header, role in ROLE_COLUMNS:
            for col in split_columns(get(row, col_header)):
                label, ix = cov.get(col, ("", ""))
                if role in ("READ", "WRITE") and not label:
                    label = "n/a"
                flat.append([
                    get(row, "Workflow"), get(row, "Workflow title"),
                    get(row, "Table"), get(row, "Op"), role, col, label, ix,
                    get(row, "Evidence"),
                ])
    return flat


def write_sheet(wb, name, headers, rows):
    ws = wb.create_sheet(title=name[:31])
    ws.append(headers)
    for r in rows:
        ws.append(r)
    head_fill = PatternFill("solid", fgColor="1F3864")
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = head_fill
        c.alignment = Alignment(vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = ws.dimensions
    for n, h in enumerate(headers, 1):
        longest = max([len(str(h))] + [len(str(r[n - 1])) for r in rows[:500] if n - 1 < len(r)])
        ws.column_dimensions[get_column_letter(n)].width = min(max(10, longest + 2), 60)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=True)
    return ws


def read_report(path):
    if os.path.isdir(path):
        files = sorted(glob.glob(os.path.join(path, "**", "*.md"), recursive=True))
        if not files:
            sys.exit(f"error: no .md files in {path}")
        parts = []
        for f in files:
            with open(f, encoding="utf-8") as fh:
                parts.append(fh.read())
        return "\n\n".join(parts)
    with open(path, encoding="utf-8") as f:
        return f.read()


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("report", help="DB_USAGE_REPORT.md or a directory of .md parts")
    ap.add_argument("--out", default="DB_USAGE_REPORT.xlsx")
    args = ap.parse_args()

    text = read_report(args.report)
    sheets = {}   # name -> [headers, rows]
    unknown = []

    for headers, rows, ctx in parse_tables(text):
        kind = classify(headers)
        if kind is None:
            unknown.append(headers)
            continue
        if kind in ("Columns", "Table_Indexes", "Foreign_Keys"):
            headers = ["Table"] + headers
            rows = [[ctx["table"]] + r for r in rows]
        elif kind == "Workflow_Access":
            headers = ["Workflow", "Workflow title"] + headers
            rows = [[ctx["workflow"], ctx["workflow_title"]] + r for r in rows]
        if kind in sheets:
            if sheets[kind][0] != headers:
                # Same kind but different headers: pad to the union, keep order.
                merged = list(sheets[kind][0])
                for h in headers:
                    if h not in merged:
                        merged.append(h)
                old_h, old_rows = sheets[kind]
                remap = lambda hs, r: [r[hs.index(h)] if h in hs else "" for h in merged]
                sheets[kind] = [merged, [remap(old_h, r) for r in old_rows] + [remap(headers, r) for r in rows]]
            else:
                sheets[kind][1].extend(rows)
        else:
            sheets[kind] = [headers, rows]

    if "Workflow_Access" in sheets:
        flat_headers = ["Workflow", "Workflow title", "Table", "Op", "Role", "Column",
                        "Index coverage", "Index", "Evidence"]
        sheets["Access_Flat"] = [flat_headers, build_flat(*sheets["Workflow_Access"])]

    if not sheets:
        sys.exit("error: no recognisable tables found. Does the report follow references/report-template.md?")

    order = ["Workflows", "Workflow_Access", "Access_Flat", "Columns", "Table_Indexes",
             "Foreign_Keys", "Index_Inventory", "Table_to_Workflows", "Column_to_Workflows", "Unresolved"]
    wb = Workbook()
    wb.remove(wb.active)
    for name in order:
        if name in sheets:
            write_sheet(wb, name, *sheets[name])
    out = os.path.abspath(args.out)
    wb.save(out)

    print(f"Wrote {out}")
    for name in order:
        if name in sheets:
            print(f"  {name:22s} {len(sheets[name][1])} rows")
    if unknown:
        print(f"\nSkipped {len(unknown)} table(s) with unrecognised headers (e.g. {unknown[0][:4]}).")
    expected = {"Workflows", "Columns", "Workflow_Access"}
    missing = expected - set(sheets)
    if missing:
        print(f"warning: no data for {sorted(missing)}; check the report follows the template headers.")


if __name__ == "__main__":
    main()
